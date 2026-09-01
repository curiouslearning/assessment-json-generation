import re
import unicodedata
import os
import codecs
import random
from datetime import datetime
import json
import csv
import openpyxl

def assessment_json_generation(sheet_id="assessment_sheet/hausa_spelling_foils_cleaned.xlsx", lang="Hausa", tab_number=0):
    assessment_type=get_assessment_bucket_title(sheet_id,tab_number)
    if "word" in assessment_type.lower():
        assessment_type= "sight-words"
    elif "spelling" in assessment_type.lower():
        assessment_type= "spelling"
    else:
        assessment_type= "letter-sounds"

    assessment_content= get_assessment_bucket(sheet_id,tab_number,assessment_type)

    if assessment_type == "spelling":
        content_version = "v0.0"
    else:
        content_version=get_and_update_content_version(sheet_id,tab_number)

    json_content = create_json_from_data(assessment_content,lang,assessment_type,content_version)
    return json_content, assessment_type

def get_assessment_bucket_title(sheet_id,tab_number):
    wb = openpyxl.load_workbook(sheet_id, data_only=True)
    worksheet = wb.worksheets[tab_number]
    tab_title = worksheet.title
    tab_name=re.sub(r"[ -]", "", tab_title)
    return tab_name

def get_and_update_content_version(sheet_id, tab_number):
    wb = openpyxl.load_workbook(sheet_id)
    worksheet = wb.worksheets[tab_number]
    
    content_version = worksheet["E1"].value
    if not content_version:
        content_version = "v0.0"
    if isinstance(content_version, str) and content_version.startswith('v'):
        numeric_version = content_version[1:]
        try:
            current_version = float(numeric_version)
            new_version = current_version + 0.1
            new_version_str = f"v{new_version:.1f}"
            worksheet["E1"] = new_version_str
            wb.save(sheet_id)
            return new_version_str
        except ValueError:
            return "Failed to parse the current version. Ensure it is in the format 'vX.Y'."
    else:
        return "Invalid or missing version format in cell E1."

def get_assessment_bucket(sheet_id,tab_number,assessment_type):
    if assessment_type == "spelling":
        max_col = 6
    elif assessment_type == "sight-words":
        max_col = 2
    else:
        max_col = 2

    wb = openpyxl.load_workbook(sheet_id, data_only=True)
    worksheet = wb.worksheets[tab_number]
    fetched_content = []
    for row in worksheet.iter_rows(min_row=2, max_row=151, min_col=1, max_col=max_col, values_only=True):
        if any(cell is not None for cell in row):
            fetched_content.append([str(c) if c is not None else "" for c in row])
    return fetched_content

def create_json_from_data(data,lang,assessment_type,content_version):
    if assessment_type == "sight-words":
        bucket_name = lang.replace(" ","-") + "sw-b"
    elif assessment_type == "spelling":
        bucket_name = lang.replace(" ","-") + "spell-b"
    else:
        bucket_name = lang.replace(" ","-") + "let-b"

    json_data = {
        "quizName": lang+" "+assessment_type.replace("-"," "),
        "appType": "assessment",
        "assessmentType": assessment_type,
        "feedbackText": "fantastic!",
        "contentVersion": content_version,
        "buckets": []
    }

    buckets = {}
    for row in data:
        if assessment_type == "spelling":
            if len(row) < 6:
                continue
        else:
            if len(row) < 2:
                continue

        bucketID = unicodedata.normalize('NFC', str(row[0]))

        if bucketID not in buckets:
            try:
                b_id_int = int(float(bucketID))
            except ValueError:
                continue

            buckets[bucketID] = {
                "bucketID": b_id_int,
                "bucketName": f"{bucket_name}-{b_id_int}",
                "usedItems": [],
                "items": []
            }

        if assessment_type == "spelling":
            item = {
                "itemName": unicodedata.normalize('NFC', str(row[1])),
                "itemText": unicodedata.normalize('NFC', str(row[2])),
                "foils": [
                    unicodedata.normalize('NFC', str(row[3])),
                    unicodedata.normalize('NFC', str(row[4])),
                    unicodedata.normalize('NFC', str(row[5]))
                ]
            }
        else:
            itemName = unicodedata.normalize('NFC', str(row[1]))
            itemText = unicodedata.normalize('NFC', str(row[1]))
            item = {
                "itemName": itemName,
                "itemText": itemText
            }

        buckets[bucketID]["items"].append(item)

    json_data["buckets"] = list(buckets.values())
    return json.dumps(json_data, indent=2, ensure_ascii=False)

if __name__ == "__main__":
    import sys
    print("=== Assessment JSON Generator ===")
    lang_input = input("Enter language (e.g., Urdu, English): ").strip()
    type_input = input("Enter type of assessment (e.g., letter-sounds, sight-words, spelling): ").strip().lower()

    sheet_dir = "assessment_sheet"
    sheet_id = None
    if os.path.exists(sheet_dir):
        for file in os.listdir(sheet_dir):
            if file.endswith(".xlsx") and lang_input.lower() in file.lower():
                sheet_id = os.path.join(sheet_dir, file)
                break
                
    if not sheet_id:
        print(f"Error: Could not find a worksheet for language '{lang_input}' in {sheet_dir}/")
        sys.exit(1)
        
    print(f"Found worksheet: {sheet_id}")
    
    try:
        wb = openpyxl.load_workbook(sheet_id, data_only=True)
    except Exception as e:
        print(f"Error loading workbook: {e}")
        sys.exit(1)
        
    tab_number = -1
    for i, sheet_name in enumerate(wb.sheetnames):
        normalized_name = sheet_name.lower()
        if type_input == "sight-words" and "word" in normalized_name:
            tab_number = i
            break
        elif type_input == "letter-sounds" and ("letter" in normalized_name or "sound" in normalized_name):
            tab_number = i
            break
        elif type_input == "spelling" and "spelling" in normalized_name:
            tab_number = i
            break

    if tab_number == -1:
        print(f"Error: Could not find a tab for '{type_input}' in the worksheet.")
        sys.exit(1)
        
    print(f"Selected tab: '{wb.sheetnames[tab_number]}' (Index: {tab_number})")
    
    json_content, assessment_type = assessment_json_generation(
        sheet_id=sheet_id,
        lang=lang_input,
        tab_number=tab_number
    )
    
    filename = f"{lang_input.lower()}-{assessment_type.replace('-', '')}.json"
    with open(filename, "w", encoding="utf-8") as f:
        f.write(json_content)
    print(f"Success! {filename} has been generated.")